import streamlit as st
from website import compiledata
from pptx import Presentation
import FastAPI

st.title("HDB Ops Survey Report Generator")

def main():
    data = compiledata()
    if data:
        createppt()
        compileppt(data)
        downloadppt()
        logusage()

def createppt():
    # Initialize a blank presentation
    prs = Presentation()

    # Add a title slide layout (usually layout index 0)
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)

    # Populate title and subtitle
    title = slide.shapes.title
    subtitle = slide.placeholders[1]

    title.text = "Hello, World!"
    subtitle.text = "Generated automatically using Python"

    # Save the presentation
    prs.save("my_presentation.pptx")


def compileppt(data):
    ...

def downloadppt():
    ...

def logusage():
    ...
    
